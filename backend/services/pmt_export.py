"""Exports des dossiers de transport vers le logiciel de facturation du client.

Il n'existe pas de format d'import commun aux logiciels de facturation des
transporteurs sanitaires (ISIS / Lomaco, Mélusine, Drivesoft, Tele Ambu, MK2i,
SCR'AMBGES…) : chacun a son propre import, souvent un tableur. Le flux B2
vers la CPAM est réservé aux logiciels certifiés CNDA, on ne le produit pas.

On propose donc :
- des formats prêts à l'emploi (CSV Excel français, CSV Windows-ANSI pour les
  logiciels anciens, XLSX, JSON, ZIP de dossiers complets avec le scan pour
  les pièces justificatives SCOR) ;
- des profils personnalisés par client : choix, ordre et intitulé des
  colonnes, séparateur, encodage, format des dates et des cases à cocher,
  pour coller au modèle d'import de son logiciel.
"""

from __future__ import annotations

import csv
import datetime as dt
import io
import json
import re
import zipfile
from typing import Any, Callable, Optional

from pydantic import BaseModel, ConfigDict, Field, field_validator

from backend.services import pmt

FORMATS = ("csv", "xlsx", "json", "zip")
ENCODINGS = {"utf-8-sig": "UTF-8 (Excel)", "utf-8": "UTF-8 sans BOM", "cp1252": "Windows-1252 (ANSI)"}
DATE_FORMATS = {"%d/%m/%Y": "JJ/MM/AAAA", "%Y-%m-%d": "AAAA-MM-JJ", "%d%m%Y": "JJMMAAAA", "%d-%m-%Y": "JJ-MM-AAAA"}
BOOL_FORMATS = {"O/N": ("O", "N"), "oui/non": ("oui", "non"), "1/0": ("1", "0"), "X/": ("X", "")}
DELIMITERS = {";": "Point-virgule", ",": "Virgule", "\t": "Tabulation", "|": "Barre verticale"}


# ──────────────────────────────────────────────────────────────────────────────
# Catalogue des champs exportables
# ──────────────────────────────────────────────────────────────────────────────

class _Ctx:
    """Données d'un dossier prêtes à être formatées."""

    def __init__(self, voucher: dict):
        self.voucher = voucher
        self.analysis = pmt.analyze(voucher.get("data"), voucher.get("transport"))
        self.d = self.analysis["data"]
        self.t = self.analysis["transport"]
        self.b = self.d["beneficiaire"]
        self.p = self.d["prescripteur"]
        self.tr = self.d["trajet"]
        self.s = self.d["situation"]
        org = re.sub(r"\D", "", self.d["organisme"]["code"])
        # Code organisme : régime (2) + caisse (3) + centre (4).
        self.regime, self.caisse, self.centre = (org[:2], org[2:5], org[5:9]) if len(org) >= 5 else ("", "", "")
        m = re.search(r"\b(\d{5})\s+([^\d,]+)$", self.b["adresse"].strip())
        self.cp, self.ville = (m.group(1), m.group(2).strip()) if m else ("", "")
        self.rue = self.b["adresse"][: m.start()].strip(" ,") if m else self.b["adresse"]


def _lieu(ctx: _Ctx, which: str) -> str:
    type_, libelle = ctx.tr[f"{which}_type"], ctx.tr[f"{which}_libelle"]
    if type_ == "domicile" and not libelle:
        return ctx.b["adresse"] or "Domicile"
    return libelle


def _exoneration(ctx: _Ctx) -> str:
    if ctx.s["at_mp"]:
        return "AT/MP"
    if ctx.s["ald_exonerante"]:
        return "ALD exonérante"
    if ctx.d["exoneration_tm"]:
        return "Exonération TM"
    if ctx.d["pension_militaire"]:
        return "Pension militaire (L. 115)"
    return ""


def _km_total(ctx: _Ctx):
    km = ctx.t["km_aller"]
    if km is None:
        return None
    return km * 2 if ctx.tr["aller_retour"] else km


MODE_CODES = {"ambulance": "AMB", "tap": "VSL/TAXI", "vehicule_personnel": "VP", "transport_commun": "TC"}

# (clé, intitulé par défaut, groupe, type, extracteur)
# type : text | date | bool | number
FIELDS: list[tuple[str, str, str, str, Callable[[_Ctx], Any]]] = [
    ("dossier_id", "Réf. dossier", "Dossier", "text", lambda c: str(c.voucher.get("id", ""))[:8]),
    ("statut", "Statut", "Dossier", "text", lambda c: c.analysis["readiness_label"]),
    ("date_transport", "Date transport", "Course", "date", lambda c: c.t["date_transport"]),
    ("heure_depart", "Heure départ", "Course", "text", lambda c: c.t["heure_depart"]),
    ("patient_nom", "Nom", "Patient", "text", lambda c: c.b["nom"]),
    ("patient_prenom", "Prénom", "Patient", "text", lambda c: c.b["prenom"]),
    ("patient_date_naissance", "Date naissance", "Patient", "date", lambda c: c.b["date_naissance"]),
    ("nir", "NIR", "Patient", "text", lambda c: c.b["nir"]),
    ("nir_cle", "Clé NIR", "Patient", "text", lambda c: c.b["nir_cle"]),
    ("nir_complet", "NIR + clé", "Patient", "text", lambda c: c.b["nir"] + c.b["nir_cle"]),
    ("patient_adresse", "Adresse", "Patient", "text", lambda c: c.rue),
    ("patient_cp", "Code postal", "Patient", "text", lambda c: c.cp),
    ("patient_ville", "Ville", "Patient", "text", lambda c: c.ville),
    ("assure_nom", "Assuré ouvrant droit", "Patient", "text", lambda c: c.d["assure"]["nom_prenom"]),
    ("assure_nir", "NIR assuré", "Patient", "text", lambda c: c.d["assure"]["nir"]),
    ("caisse_libelle", "Caisse", "Droits", "text", lambda c: c.d["organisme"]["libelle"]),
    ("regime", "Code régime", "Droits", "text", lambda c: c.regime),
    ("caisse", "Code caisse", "Droits", "text", lambda c: c.caisse),
    ("centre", "Code centre", "Droits", "text", lambda c: c.centre),
    ("nature_assurance", "Nature assurance", "Droits", "text", lambda c: "AT/MP" if c.s["at_mp"] else "Maladie"),
    ("date_at_mp", "Date AT/MP", "Droits", "date", lambda c: c.s["date_at_mp"]),
    ("hospitalisation", "Hospitalisation", "Droits", "bool", lambda c: c.s["hospitalisation"]),
    ("ald", "ALD", "Droits", "bool", lambda c: c.s["ald_exonerante"] or c.s["ald_non_exonerante"]),
    ("exoneration", "Exonération", "Droits", "text", _exoneration),
    ("taux_amo", "Taux AMO", "Droits", "number", lambda c: c.analysis["prise_en_charge"]["taux_amo"]),
    ("accident_tiers", "Accident tiers", "Droits", "bool", lambda c: c.d["accident_tiers"]),
    ("date_accident", "Date accident", "Droits", "date", lambda c: c.d["date_accident"]),
    ("mode_code", "Code mode", "Transport", "text", lambda c: MODE_CODES.get(c.d["mode"] or "", "")),
    ("mode", "Mode", "Transport", "text", lambda c: pmt.MODES.get(c.d["mode"] or "", "")),
    ("justification", "Justification ambulance", "Transport", "text",
     lambda c: ", ".join(v for k, v in pmt.AMBULANCE_JUSTIFS.items() if c.d["ambulance_justif"].get(k))),
    ("transport_partage", "Transport partagé", "Transport", "bool", lambda c: c.d["transport_partage"]),
    ("urgence", "Urgence", "Transport", "bool", lambda c: c.d["urgence"]["samu"] or c.d["urgence"]["autre"]),
    ("depart", "Départ", "Trajet", "text", lambda c: _lieu(c, "depart")),
    ("arrivee", "Arrivée", "Trajet", "text", lambda c: _lieu(c, "arrivee")),
    ("aller_retour", "Aller-retour", "Trajet", "bool", lambda c: c.tr["aller_retour"]),
    ("nb_iteratifs", "Transports itératifs", "Trajet", "number", lambda c: c.tr["nb_iteratifs"]),
    ("km_aller", "Km aller", "Trajet", "number", lambda c: c.t["km_aller"]),
    ("km_total", "Km total", "Trajet", "number", _km_total),
    ("prescripteur", "Prescripteur", "Prescription", "text", lambda c: c.p["nom"]),
    ("rpps", "RPPS", "Prescription", "text", lambda c: c.p["rpps"]),
    ("structure", "Structure", "Prescription", "text", lambda c: c.p["raison_sociale"]),
    ("structure_numero", "N° structure", "Prescription", "text", lambda c: c.p["numero_structure"]),
    ("date_prescription", "Date prescription", "Prescription", "date", lambda c: c.p["date_prescription"]),
    ("accord_prealable", "Accord préalable", "Course", "text", lambda c: c.t["accord_prealable_ref"]),
    ("vehicule", "Véhicule", "Course", "text", lambda c: c.t["vehicule"]),
    ("equipage", "Équipage", "Course", "text", lambda c: c.t["equipage"]),
    ("nb_erreurs", "Erreurs", "Contrôles", "number", lambda c: c.analysis["counts"]["error"]),
    ("nb_alertes", "Alertes", "Contrôles", "number", lambda c: c.analysis["counts"]["warning"]),
    ("anomalies", "Anomalies", "Contrôles", "text",
     lambda c: " | ".join(x["message"] for x in c.analysis["checks"] if x["level"] in ("error", "warning"))),
    ("fichier_scan", "Fichier PMT", "Dossier", "text", lambda c: c.voucher.get("scan_name", "")),
]
FIELD_INDEX = {f[0]: f for f in FIELDS}

STANDARD_COLUMNS = [
    "dossier_id", "date_transport", "heure_depart", "patient_nom", "patient_prenom", "patient_date_naissance",
    "nir", "nir_cle", "patient_adresse", "patient_cp", "patient_ville", "caisse_libelle", "regime", "caisse",
    "centre", "nature_assurance", "exoneration", "taux_amo", "accident_tiers", "mode_code", "justification",
    "depart", "arrivee", "aller_retour", "nb_iteratifs", "km_aller", "km_total", "prescripteur", "rpps",
    "structure_numero", "date_prescription", "accord_prealable", "vehicule", "equipage", "statut", "anomalies",
]


def field_catalog() -> list[dict]:
    return [{"key": k, "label": label, "group": g, "type": t} for k, label, g, t, _ in FIELDS]


# ──────────────────────────────────────────────────────────────────────────────
# Profils d'export
# ──────────────────────────────────────────────────────────────────────────────

class Column(BaseModel):
    model_config = ConfigDict(extra="ignore")
    key: str
    label: str = ""

    @field_validator("key")
    @classmethod
    def _known(cls, v: str) -> str:
        if v not in FIELD_INDEX:
            raise ValueError(f"Champ inconnu : {v}")
        return v


class ExportProfile(BaseModel):
    model_config = ConfigDict(extra="ignore")
    name: str = "Export"
    format: str = "csv"
    delimiter: str = ";"
    encoding: str = "utf-8-sig"
    date_format: str = "%d/%m/%Y"
    bool_format: str = "O/N"
    decimal: str = ","
    header: bool = True
    columns: list[Column] = Field(default_factory=lambda: [Column(key=k) for k in STANDARD_COLUMNS])

    @field_validator("format")
    @classmethod
    def _fmt(cls, v):
        if v not in FORMATS:
            raise ValueError(f"Format inconnu : {v}")
        return v

    @field_validator("delimiter")
    @classmethod
    def _delim(cls, v):
        if v not in DELIMITERS:
            raise ValueError("Séparateur non pris en charge")
        return v

    @field_validator("encoding")
    @classmethod
    def _enc(cls, v):
        if v not in ENCODINGS:
            raise ValueError("Encodage non pris en charge")
        return v

    @field_validator("date_format")
    @classmethod
    def _date(cls, v):
        if v not in DATE_FORMATS:
            raise ValueError("Format de date non pris en charge")
        return v

    @field_validator("bool_format")
    @classmethod
    def _bool(cls, v):
        if v not in BOOL_FORMATS:
            raise ValueError("Format oui/non non pris en charge")
        return v

    @field_validator("decimal")
    @classmethod
    def _dec(cls, v):
        if v not in (",", "."):
            raise ValueError("Séparateur décimal : « , » ou « . »")
        return v

    @field_validator("columns")
    @classmethod
    def _cols(cls, v):
        if not v:
            raise ValueError("Choisir au moins une colonne")
        return v


PRESETS: dict[str, dict] = {
    "standard": {"name": "CSV Excel (recommandé)", "format": "csv"},
    "ansi": {"name": "CSV Windows-ANSI (logiciels anciens)", "format": "csv", "encoding": "cp1252"},
    "xlsx": {"name": "Classeur Excel (.xlsx)", "format": "xlsx"},
    "json": {"name": "JSON (intégration API)", "format": "json", "date_format": "%Y-%m-%d", "decimal": ".",
             "bool_format": "1/0"},
    "zip": {"name": "Dossiers complets (ZIP : tableau + PMT scannées + fiches)", "format": "zip"},
}


def preset(name: str) -> ExportProfile:
    if name not in PRESETS:
        raise ValueError(f"Format prédéfini inconnu : {name}")
    return ExportProfile.model_validate(PRESETS[name])


def preset_catalog() -> list[dict]:
    return [{"id": k, **ExportProfile.model_validate(v).model_dump()} for k, v in PRESETS.items()]


def options_catalog() -> dict:
    return {
        "formats": list(FORMATS),
        "encodings": ENCODINGS,
        "date_formats": DATE_FORMATS,
        "bool_formats": list(BOOL_FORMATS),
        "delimiters": DELIMITERS,
    }


# ──────────────────────────────────────────────────────────────────────────────
# Mise en forme
# ──────────────────────────────────────────────────────────────────────────────

def _format(value: Any, type_: str, profile: ExportProfile) -> Any:
    if type_ == "bool":
        if value is None:
            return ""
        yes, no = BOOL_FORMATS[profile.bool_format]
        return yes if value else no
    if type_ == "date":
        d = pmt.parse_date(value)
        return d.strftime(profile.date_format) if d else (value or "")
    if type_ == "number":
        if value is None or value == "":
            return ""
        if isinstance(value, float) and value.is_integer():
            value = int(value)
        return str(value).replace(".", profile.decimal) if profile.format in ("csv", "zip") else value
    return "" if value is None else str(value)


def build_rows(vouchers: list[dict], profile: ExportProfile) -> tuple[list[str], list[list[Any]]]:
    headers = [c.label or FIELD_INDEX[c.key][1] for c in profile.columns]
    rows = []
    for v in vouchers:
        ctx = _Ctx(v)
        rows.append([_format(FIELD_INDEX[c.key][4](ctx), FIELD_INDEX[c.key][3], profile) for c in profile.columns])
    return headers, rows


def to_csv(vouchers: list[dict], profile: ExportProfile) -> bytes:
    headers, rows = build_rows(vouchers, profile)
    buf = io.StringIO()
    writer = csv.writer(buf, delimiter=profile.delimiter, lineterminator="\r\n")
    if profile.header:
        writer.writerow(headers)
    writer.writerows(rows)
    return buf.getvalue().encode(profile.encoding, errors="replace")


def to_xlsx(vouchers: list[dict], profile: ExportProfile) -> bytes:
    from openpyxl import Workbook
    from openpyxl.styles import Font

    headers, rows = build_rows(vouchers, profile)
    wb = Workbook()
    ws = wb.active
    ws.title = "Transports"
    if profile.header:
        ws.append(headers)
        for cell in ws[1]:
            cell.font = Font(bold=True)
        ws.freeze_panes = "A2"
    for row in rows:
        ws.append(row)
    # Les identifiants (NIR, RPPS, codes caisse) restent du texte : pas de 1,83E+12.
    for col_idx, col in enumerate(profile.columns, start=1):
        if FIELD_INDEX[col.key][3] == "text":
            for cell in ws.iter_cols(min_col=col_idx, max_col=col_idx, min_row=2):
                for c in cell:
                    c.number_format = "@"
        ws.column_dimensions[ws.cell(row=1, column=col_idx).column_letter].width = max(
            12, min(40, len(headers[col_idx - 1]) + 4))
    out = io.BytesIO()
    wb.save(out)
    return out.getvalue()


def to_json(vouchers: list[dict], profile: ExportProfile) -> bytes:
    headers, rows = build_rows(vouchers, profile)
    keys = [c.key for c in profile.columns]
    items = []
    for v, row in zip(vouchers, rows):
        item = dict(zip(keys, row))
        item["_dossier"] = {"id": v.get("id"), "data": v.get("data"), "transport": v.get("transport")}
        items.append(item)
    payload = {
        "format": "pulse-pme.transports.v1",
        "genere_le": dt.datetime.now(dt.timezone.utc).isoformat(),
        "nombre": len(items),
        "dossiers": items,
    }
    return json.dumps(payload, ensure_ascii=False, indent=2).encode("utf-8")


def _slug(text: str) -> str:
    text = re.sub(r"[^A-Za-z0-9]+", "-", text.upper()).strip("-")
    return text[:40] or "PATIENT"


def dossier_name(v: dict) -> str:
    """AAAAMMJJ_NOM_PRENOM_ref : tri chronologique et rattachement facile au bon patient."""
    analysis_data = pmt.normalize_pmt(v.get("data"))
    t = pmt.normalize_transport(v.get("transport"))
    d = pmt.parse_date(t.date_transport) or pmt.parse_date(analysis_data.prescripteur.date_prescription)
    b = analysis_data.beneficiaire
    return f"{d.strftime('%Y%m%d') if d else 'SANS-DATE'}_{_slug(b.nom)}_{_slug(b.prenom)}_{str(v.get('id', ''))[:8]}"


SCAN_EXT = {"application/pdf": ".pdf", "image/jpeg": ".jpg", "image/png": ".png", "image/webp": ".webp",
            "image/heic": ".heic"}


def to_zip(vouchers: list[dict], profile: ExportProfile, read_scan: Callable[[dict], Optional[bytes]]) -> bytes:
    """Un dossier par transport : PMT scannée (pièce justificative SCOR), fiche, données ; tableau récapitulatif."""
    out = io.BytesIO()
    named = []
    with zipfile.ZipFile(out, "w", zipfile.ZIP_DEFLATED) as zf:
        for v in vouchers:
            folder = dossier_name(v)
            item = dict(v)
            scan = read_scan(v)
            if scan:
                scan_name = f"{folder}/PMT_{folder}{SCAN_EXT.get(v.get('scan_mime') or '', '.bin')}"
                zf.writestr(scan_name, scan)
                item["scan_name"] = scan_name
            zf.writestr(f"{folder}/fiche.html", pmt.render_fiche_html(v))
            zf.writestr(f"{folder}/dossier.json", json.dumps(
                {"id": v.get("id"), "data": v.get("data"), "transport": v.get("transport"),
                 "transporteur": v.get("transporteur")}, ensure_ascii=False, indent=2))
            named.append(item)
        csv_profile = profile.model_copy(update={"format": "csv"})
        if not any(c.key == "fichier_scan" for c in csv_profile.columns):
            csv_profile.columns = [*csv_profile.columns, Column(key="fichier_scan")]
        zf.writestr("recapitulatif.csv", to_csv(named, csv_profile))
        zf.writestr("LISEZMOI.txt", (
            "Export des dossiers de transport sanitaire\r\n\r\n"
            "recapitulatif.csv : une ligne par transport, à importer dans le logiciel de facturation.\r\n"
            "Chaque dossier contient la PMT scannée (pièce justificative à joindre via SCOR),\r\n"
            "une fiche imprimable et les données au format JSON.\r\n\r\n"
            "Données de santé : à conserver sur un poste sécurisé et à supprimer après facturation.\r\n"
        ).encode("utf-8"))
    return out.getvalue()


MEDIA = {
    "csv": ("text/csv", "csv"),
    "xlsx": ("application/vnd.openxmlformats-officedocument.spreadsheetml.sheet", "xlsx"),
    "json": ("application/json", "json"),
    "zip": ("application/zip", "zip"),
}


def render(vouchers: list[dict], profile: ExportProfile,
           read_scan: Optional[Callable[[dict], Optional[bytes]]] = None) -> tuple[bytes, str, str]:
    """Retourne (contenu, type MIME, extension)."""
    if profile.format == "csv":
        content = to_csv(vouchers, profile)
    elif profile.format == "xlsx":
        content = to_xlsx(vouchers, profile)
    elif profile.format == "json":
        content = to_json(vouchers, profile)
    else:
        content = to_zip(vouchers, profile, read_scan or (lambda v: None))
    mime, ext = MEDIA[profile.format]
    if profile.format == "csv":
        mime += f"; charset={'windows-1252' if profile.encoding == 'cp1252' else 'utf-8'}"
    return content, mime, ext


# ──────────────────────────────────────────────────────────────────────────────
# Sélection et indicateurs
# ──────────────────────────────────────────────────────────────────────────────

def is_exportable(readiness: str, status: str) -> bool:
    """Prêt sans intervention, ou « à vérifier » validé par un humain. Jamais bloquant."""
    if readiness == "bloquant":
        return False
    return readiness == "pret" or status in ("validated", "exported", "billed")


def stats(vouchers: list[dict]) -> dict:
    total = len(vouchers)
    by_readiness = {k: 0 for k in pmt.READINESS_LABELS}
    causes: dict[str, dict] = {}
    exportable = auto_ready = exported = 0
    for v in vouchers:
        r = v.get("readiness") or "a_verifier"
        by_readiness[r] = by_readiness.get(r, 0) + 1
        if r == "pret":
            auto_ready += 1
        if is_exportable(r, v.get("status") or "draft"):
            exportable += 1
        if v.get("exported_at"):
            exported += 1
        for c in v.get("checks") or []:
            if c.get("level") in ("error", "warning"):
                entry = causes.setdefault(c["code"], {"code": c["code"], "level": c["level"],
                                                      "message": c["message"], "count": 0})
                entry["count"] += 1
    pct = lambda n: round(100 * n / total, 1) if total else 0.0  # noqa: E731
    return {
        "total": total,
        "by_readiness": by_readiness,
        "taux_prets_auto": pct(auto_ready),      # prêts sans aucune intervention
        "taux_exportables": pct(exportable),     # prêts + vérifiés par un humain
        "exportes": exported,
        "a_exporter": sum(1 for v in vouchers
                          if is_exportable(v.get("readiness") or "", v.get("status") or "draft")
                          and not v.get("exported_at")),
        "principales_causes": sorted(causes.values(), key=lambda x: -x["count"])[:8],
    }
