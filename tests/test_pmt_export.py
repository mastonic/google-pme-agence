import csv
import io
import json
import os
import tempfile
import unittest
import zipfile

from backend.services import pmt_export
from tests.test_pmt import NIR, NIR_CLE, TRANSPORT, sample_pmt


def voucher(**overrides):
    return {"id": "abcdef12-0000", "data": sample_pmt(**overrides), "transport": TRANSPORT,
            "transporteur": {"raison_sociale": "Ambulances Test"}}


def parse_csv(content: bytes, encoding="utf-8-sig", delimiter=";"):
    return list(csv.reader(io.StringIO(content.decode(encoding)), delimiter=delimiter))


class FieldTests(unittest.TestCase):
    def test_standard_csv_is_excel_friendly_and_splits_codes(self):
        content = pmt_export.to_csv([voucher()], pmt_export.preset("standard"))
        self.assertTrue(content.startswith("﻿".encode("utf-8")))
        header, row = parse_csv(content)
        line = dict(zip(header, row))
        self.assertEqual(line["Date transport"], "07/10/2026")
        self.assertEqual(line["NIR"], NIR)
        self.assertEqual(line["Clé NIR"], NIR_CLE)
        self.assertEqual((line["Code régime"], line["Code caisse"], line["Code centre"]), ("01", "751", "0000"))
        self.assertEqual((line["Code postal"], line["Ville"]), ("75001", "Paris"))
        self.assertEqual(line["Code mode"], "AMB")
        self.assertEqual(line["Aller-retour"], "O")
        self.assertEqual(line["Km total"], "64")
        self.assertEqual(line["Exonération"], "ALD exonérante")

    def test_home_departure_uses_patient_address(self):
        _, row = parse_csv(pmt_export.to_csv([voucher()], pmt_export.ExportProfile(
            columns=[{"key": "depart"}])))
        self.assertEqual(row, ["1 rue de Test 75001 Paris"])

    def test_custom_profile_controls_layout(self):
        profile = pmt_export.ExportProfile(
            delimiter="\t", encoding="cp1252", date_format="%Y-%m-%d", bool_format="1/0", decimal=".",
            header=False,
            columns=[{"key": "patient_nom", "label": "NOM_PAT"}, {"key": "date_transport"},
                     {"key": "aller_retour"}, {"key": "km_aller"}, {"key": "structure"}],
        )
        content = pmt_export.to_csv([voucher(**{"prescripteur.raison_sociale": "Hôpital Évry",
                                               "trajet.aller_retour": False})], profile)
        rows = parse_csv(content, encoding="cp1252", delimiter="\t")
        self.assertEqual(rows, [["MARTIN", "2026-10-07", "0", "32", "Hôpital Évry"]])

    def test_decimal_comma(self):
        v = voucher()
        v["transport"] = {**TRANSPORT, "km_aller": 12.5}
        _, row = parse_csv(pmt_export.to_csv([v], pmt_export.ExportProfile(columns=[{"key": "km_aller"}])))
        self.assertEqual(row, ["12,5"])

    def test_profile_rejects_unknown_field_and_options(self):
        with self.assertRaises(ValueError):
            pmt_export.ExportProfile(columns=[{"key": "inconnu"}])
        with self.assertRaises(ValueError):
            pmt_export.ExportProfile(encoding="latin-9")
        with self.assertRaises(ValueError):
            pmt_export.ExportProfile(columns=[])


class FormatTests(unittest.TestCase):
    def test_xlsx_keeps_identifiers_as_text(self):
        from openpyxl import load_workbook
        content = pmt_export.to_xlsx([voucher()], pmt_export.preset("xlsx"))
        ws = load_workbook(io.BytesIO(content)).active
        header = [c.value for c in ws[1]]
        nir_cell = ws.cell(row=2, column=header.index("NIR") + 1)
        self.assertEqual(nir_cell.value, NIR)
        self.assertEqual(nir_cell.number_format, "@")

    def test_json_contains_full_dossier(self):
        payload = json.loads(pmt_export.to_json([voucher()], pmt_export.preset("json")))
        self.assertEqual(payload["nombre"], 1)
        item = payload["dossiers"][0]
        self.assertEqual(item["date_transport"], "2026-10-07")
        self.assertEqual(item["aller_retour"], "1")
        self.assertEqual(item["_dossier"]["data"]["beneficiaire"]["nir"], NIR)

    def test_zip_contains_scan_fiche_and_summary(self):
        v = {**voucher(), "scan_mime": "application/pdf"}
        content = pmt_export.to_zip([v], pmt_export.preset("zip"), read_scan=lambda _: b"%PDF-1.4 test")
        zf = zipfile.ZipFile(io.BytesIO(content))
        names = zf.namelist()
        folder = "20261007_MARTIN_CLAIRE_abcdef12"
        self.assertIn(f"{folder}/PMT_{folder}.pdf", names)
        self.assertIn(f"{folder}/fiche.html", names)
        self.assertIn(f"{folder}/dossier.json", names)
        header, row = parse_csv(zf.read("recapitulatif.csv"))
        self.assertEqual(dict(zip(header, row))["Fichier PMT"], f"{folder}/PMT_{folder}.pdf")


class SelectionTests(unittest.TestCase):
    def test_exportable_rules(self):
        self.assertTrue(pmt_export.is_exportable("pret", "draft"))
        self.assertFalse(pmt_export.is_exportable("a_verifier", "draft"))
        self.assertTrue(pmt_export.is_exportable("a_verifier", "validated"))
        self.assertFalse(pmt_export.is_exportable("bloquant", "validated"))

    def test_stats(self):
        rows = [
            {"readiness": "pret", "status": "draft", "checks": []},
            {"readiness": "pret", "status": "exported", "checks": [], "exported_at": "x"},
            {"readiness": "a_verifier", "status": "validated", "checks": [
                {"level": "warning", "code": "TIERS_NON_RENSEIGNE", "message": "m"}]},
            {"readiness": "bloquant", "status": "draft", "checks": [
                {"level": "error", "code": "SIGNATURE", "message": "s"}]},
        ]
        s = pmt_export.stats(rows)
        self.assertEqual(s["taux_prets_auto"], 50.0)
        self.assertEqual(s["taux_exportables"], 75.0)
        self.assertEqual(s["a_exporter"], 2)
        self.assertEqual({c["code"] for c in s["principales_causes"]}, {"TIERS_NON_RENSEIGNE", "SIGNATURE"})


class ExportApiTests(unittest.TestCase):
    def setUp(self):
        from fastapi import FastAPI
        from fastapi.testclient import TestClient
        from sqlalchemy import create_engine
        from sqlalchemy.orm import sessionmaker
        from sqlalchemy.pool import StaticPool

        from backend.models.database import Base, get_db
        from backend.routers.pmt import router

        self.tmp = tempfile.TemporaryDirectory()
        os.environ["PMT_SCAN_DIR"] = self.tmp.name
        engine = create_engine("sqlite://", connect_args={"check_same_thread": False}, poolclass=StaticPool)
        Base.metadata.create_all(bind=engine)
        Session = sessionmaker(bind=engine)

        def override():
            db = Session()
            try:
                yield db
            finally:
                db.close()

        app = FastAPI()
        app.include_router(router)
        app.dependency_overrides[get_db] = override
        from backend.routers.pmt_auth import get_current_user
        from backend.services.pmt_auth import CurrentUser
        app.dependency_overrides[get_current_user] = lambda: CurrentUser(1, "admin@test.fr", "admin", None)
        self.client = TestClient(app)

    def tearDown(self):
        os.environ.pop("PMT_SCAN_DIR", None)
        self.tmp.cleanup()

    def create(self, **overrides):
        r = self.client.post("/pmt/vouchers", json={"business_id": "amb-1", "data": sample_pmt(**overrides),
                                                    "transport": TRANSPORT})
        self.assertEqual(r.status_code, 200)
        return r.json()

    def test_only_ready_dossiers_are_exported_once(self):
        ready = self.create()
        self.create(**{"prescripteur.signature_presente": False})   # bloquant
        review = self.create(accident_tiers=None)                    # à vérifier

        stats = self.client.get("/pmt/stats", params={"business_id": "amb-1"}).json()
        self.assertEqual(stats["total"], 3)
        self.assertEqual(stats["a_exporter"], 1)

        r = self.client.post("/pmt/export", json={"business_id": "amb-1", "preset": "standard"})
        self.assertEqual(r.status_code, 200)
        self.assertEqual(r.headers["x-export-count"], "1")
        self.assertEqual(len(parse_csv(r.content)), 2)
        self.assertEqual(self.client.get(f"/pmt/vouchers/{ready['id']}").json()["status"], "exported")

        # Déjà exporté : pas de doublon au prochain export.
        self.assertEqual(self.client.post("/pmt/export", json={"business_id": "amb-1"}).status_code, 404)

        # Le dossier « à vérifier » part une fois validé par un humain.
        self.client.patch(f"/pmt/vouchers/{review['id']}", json={"status": "validated"})
        r = self.client.post("/pmt/export", json={"business_id": "amb-1", "preset": "xlsx"})
        self.assertEqual(r.headers["x-export-count"], "1")
        self.assertTrue(r.content.startswith(b"PK"))

    def test_editing_exported_dossier_requires_reexport(self):
        v = self.create()
        self.client.post("/pmt/export", json={"business_id": "amb-1"})
        r = self.client.patch(f"/pmt/vouchers/{v['id']}", json={"transport": {**TRANSPORT, "km_aller": 40}})
        self.assertEqual(r.json()["status"], "draft")

    def test_custom_profile_crud_and_export(self):
        self.create()
        r = self.client.post("/pmt/export/profiles", json={
            "business_id": "amb-1", "name": "Import logiciel X", "delimiter": "|",
            "columns": [{"key": "nir_complet", "label": "NUMSS"}, {"key": "date_transport", "label": "DATE"}]})
        self.assertEqual(r.status_code, 200)
        pid = r.json()["id"]
        self.assertEqual(len(self.client.get("/pmt/export/profiles", params={"business_id": "amb-1"}).json()), 1)
        r = self.client.post("/pmt/export", json={"business_id": "amb-1", "profile_id": pid})
        self.assertEqual(parse_csv(r.content, delimiter="|"), [["NUMSS", "DATE"], [NIR + NIR_CLE, "07/10/2026"]])
        bad = self.client.post("/pmt/export/profiles", json={"name": "x", "columns": [{"key": "nope"}]})
        self.assertEqual(bad.status_code, 400)
        self.assertEqual(self.client.delete(f"/pmt/export/profiles/{pid}").status_code, 200)

    def test_scan_is_kept_for_zip_and_deleted_with_dossier(self):
        from unittest import mock
        from backend.services import pmt

        fake = {"data": pmt.normalize_pmt(sample_pmt()), "provider": "test"}
        with mock.patch.object(pmt, "extract_pmt", return_value=fake):
            r = self.client.post("/pmt/extract", data={"business_id": "amb-1"},
                                 files={"file": ("bon.pdf", b"%PDF-1.4 scan", "application/pdf")})
        self.assertEqual(r.status_code, 200)
        v = r.json()
        self.assertTrue(v["has_scan"])
        self.client.patch(f"/pmt/vouchers/{v['id']}", json={"transport": TRANSPORT})
        self.assertEqual(self.client.get(f"/pmt/vouchers/{v['id']}/scan").content, b"%PDF-1.4 scan")

        r = self.client.post("/pmt/export", json={"business_id": "amb-1", "preset": "zip"})
        zf = zipfile.ZipFile(io.BytesIO(r.content))
        scans = [n for n in zf.namelist() if "/PMT_" in n]
        self.assertEqual(len(scans), 1)
        self.assertEqual(zf.read(scans[0]), b"%PDF-1.4 scan")

        self.client.delete(f"/pmt/vouchers/{v['id']}")
        self.assertEqual(os.listdir(self.tmp.name), [])

    def test_options_catalog(self):
        opts = self.client.get("/pmt/export/options").json()
        self.assertIn("nir", {f["key"] for f in opts["fields"]})
        self.assertEqual({p["id"] for p in opts["presets"]}, {"standard", "ansi", "xlsx", "json", "zip"})


if __name__ == "__main__":
    unittest.main()
